"""Excel de prueba FICTICIO con la misma estructura y problemas que el original.

Problemas incluidos: unidades en las celdas ("65kg", "3.10s"), '-' por ausencia, filas de
metadatos de hora, hoja «hoy» sin título, posición «opuesto-central», talla inconsistente,
jugador sin consentimiento, atípico en salto, nombres con distinto uso de mayúsculas/espacios,
velocidad mal digitada y título que contradice el nombre de la hoja.
"""

from pathlib import Path

from openpyxl import Workbook

NOMBRES = {
    "A": "Alfa Prueba Carlos Eduardo",
    "B": "Beta Prueba Diego",
    "C": "Gamma Prueba Mario Luis",  # opuesto-central, talla inconsistente
    "D": "Delta Prueba Juan",  # sin consentimiento
    "E": "Epsilon Prueba Ruben Mateo",
    "F": "Zeta Prueba Hugo",
    "G": "Eta Prueba Pedro Nel",
    "H": "Theta Prueba Luis",  # salto atípico
}
EDAD = {"A": 22, "B": 24, "C": 24, "D": 23, "E": 19, "F": 18, "G": 20, "H": 21}
POSICION = {
    "A": "punta receptora",
    "B": "armador ",
    "C": "opuesto-central",
    "D": "libero",
    "E": "central",
    "F": "opuesto",
    "G": "punta receptora",
    "H": "central",
}
PESO = {"A": 65, "B": 79, "C": 101, "D": 56, "E": 71, "F": 72, "G": 67, "H": 76}

COLS_SALTO = [
    "Nombre del deportista",
    "edad",
    "talla",
    "peso (kg)",
    "salto sin previo calentamiento",
    "salto despues de un breve calentamiento",
    "salto despues de una activacion intensa",
    "salto despues de actividad fisica intensa ",
    "Posiicion de campo",
]
COLS_VEL = [
    "Nombre del deportista",
    "edad",
    "talla",
    "peso (kg)",
    "Sprint 20m (tiempo)",
    "Velocidad (m/s)",
    "Palier Course Navette",
    "VO2max estimado (ml/kg/min)",
    "Tiempo de entrenamiento efectivo (min)",
    "Posiicion de campo",
]


def _talla(k: str, hoja_idx: int) -> float:
    base = {"A": 1.78, "B": 1.82, "C": 1.88, "D": 1.69, "E": 1.79, "F": 1.79, "G": 1.77, "H": 1.85}
    if k == "C":
        return [1.88, 1.87, 1.86, 1.87][hoja_idx % 4]
    return base[k]


def _hoja_salto(wb, nombre, titulo, idx, ausente=None, atipico=True):
    ws = wb.create_sheet(nombre)
    ws.append([titulo] + [None] * 8)
    ws.append(COLS_SALTO)
    for n, k in enumerate("ABCDEFGH"):
        base = 20 + n
        if k == ausente:
            vals = ["-", "-", "-", "-"]
        else:
            vals = [f"{base}cm", f"{base + 7}cm", f"{base + 12}cm", f"{base + 5}cm"]
        if k == "H" and atipico:
            vals = ["70cm", "75cm", "80cm", "72cm"]
        nombre_excel = NOMBRES[k]
        if idx % 2:  # variantes de escritura entre hojas
            nombre_excel = nombre_excel.upper() + " "
        ws.append([nombre_excel, EDAD[k], _talla(k, idx), f"{PESO[k]}kg", *vals, POSICION[k]])
    ws.append(
        [
            "HORA DE LA TOMA DE LOS DATOS",
            None,
            None,
            None,
            "7:00PM a 7:15 PM",
            "7:30 a 7:45",
            "8:30 a 8:45",
            "9:00 a 9:15",
            None,
        ]
    )


def _hoja_vel(wb, nombre, titulo, idx, vel_mala=False):
    ws = wb.create_sheet(nombre)
    ws.append([titulo] + [None] * 9)
    ws.append(COLS_VEL)
    for n, k in enumerate("ABCDEFGH"):
        t = round(3.00 + 0.05 * n + 0.01 * idx, 2)
        palier = 9 + 0.5 * n
        edad = EDAD[k]
        v = 7.5 + 0.5 * palier
        vo2 = round(31.025 + 3.238 * v - 3.248 * edad + 0.1536 * edad * v, 1)
        vel = round(20 / t, 2)
        if vel_mala and k == "B":
            vel = round(vel + 0.3, 2)
        fila = [
            NOMBRES[k],
            edad,
            _talla(k, idx),
            f"{PESO[k]}kg",
            f"{t:.2f}s",
            f"{vel:.2f}m/s",
            palier,
            f"{vo2:.1f}",
            f"{70 + n}min",
            POSICION[k].strip(),
        ]
        if k == "E" and idx == 2:  # ausencia
            fila[4:9] = ["-"] * 5
        ws.append(fila)
    ws.append(["HORA - SPRINT 20m", None, None, None, "6:45 PM a 7:00 PM"] + [None] * 5)
    ws.append(
        ["HORA - COURSE NAVETTE", None, None, None, None, None, "7:15 PM a 7:45 PM"] + [None] * 3
    )
    ws.append(
        [
            "HORA - REGISTRO ENTRENAMIENTO",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            "5:30 PM a 7:00 PM",
            None,
        ]
    )


def crear_excel_ficticio(ruta: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    _hoja_salto(wb, "10 Agosto", "PRUEBA DE SALTO VERTICAL LUNES 10 DE AGOSTO DEL 2026", 0)
    _hoja_salto(
        wb, "12 Agosto", "PRUEBA DE SALTO VERTICAL MIERCOLES 12 DE AGOSTO DEL 2026", 1, ausente="F"
    )
    _hoja_vel(wb, "14 Agosto", "PRUEBAS DE VELOCIDAD - VIERNES 14 DE AGOSTO DEL 2026", 0)
    # «hoy»: sin título (fila 1 vacía) → exige --fecha-hoy
    ws = wb.create_sheet("hoy")
    _hoja_vel(wb, "tmp", "x", 1)
    tmp = wb["tmp"]
    ws.append([None] * 10)
    for fila in list(tmp.iter_rows(min_row=2, values_only=True)):
        ws.append(list(fila))
    wb.remove(tmp)
    # título (martes 18) contradice el nombre de la hoja (17): se usa el título
    _hoja_vel(
        wb, "17 Agosto", "PRUEBAS DE VELOCIDAD - MARTES 18 DE AGOSTO DEL 2026", 2, vel_mala=True
    )
    wb.save(ruta)
    return ruta


CONSENTIMIENTOS_CSV = """nombre_normalizado,consentimiento_informado,fecha_consentimiento
Alfa Prueba Carlos Eduardo,si,2026-09-01
Prueba Beta Diego,si,2026-09-01
Gamma Prueba Mario Luis,si,2026-09-01
Epsilon Prueba Ruben Mateo,si,2026-09-01
Zeta Prueba Hugo,si,2026-09-01
Eta Prueba Pedro Nel,si,2026-09-01
Theta Prueba Luis,si,2026-09-01
Nadie Que Existe,si,2026-09-01
"""

DECISIONES_CSV = """tipo,clave,valor
posicion,opuesto-central,opuesto
nombre,Epsilon Prueba Ruben Mateo,Ruben Mateo|Epsilon Prueba
"""
